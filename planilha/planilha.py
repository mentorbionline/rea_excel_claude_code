"""Planilha — editor de planilhas simples, escrito do zero (tkinter + openpyxl).

Modelo puro (células, fórmulas, formatos, desfazer/refazer, xlsx) no topo;
interface tkinter na classe App, no final.
"""
import os
import re
import sys
import warnings
from collections import namedtuple
from copy import copy
from datetime import date, datetime, time, timedelta
from functools import lru_cache

import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog, colorchooser
import tkinter.font as tkfont

# ponytail: avaliação recursiva; cadeias de dependência muito longas (>~3000 células) pedem ordenação topológica iterativa
sys.setrecursionlimit(20000)

# ---------------------------------------------------------------- modelo
FUNCS = {'SOMA': 'SUM', 'SUM': 'SUM',
         'MÉDIA': 'AVERAGE', 'MEDIA': 'AVERAGE', 'AVERAGE': 'AVERAGE',
         'MÍNIMO': 'MIN', 'MINIMO': 'MIN', 'MIN': 'MIN',
         'MÁXIMO': 'MAX', 'MAXIMO': 'MAX', 'MAX': 'MAX'}
FORMATS = ['Geral', 'Número', 'Moeda', 'Porcentagem', 'Data']
# formato da célula; code = código de formato original do arquivo (preservado enquanto nf/dec não mudam)
Fmt = namedtuple('Fmt', 'bold italic nf dec color fill align wrap valign code')
DEF = Fmt(False, False, 'Geral', 2, None, None, None, False, None, None)
DATE_CODE = 'dd/mm/yyyy'


def nf_code(nf, dec):
    """Código de formato do Excel gerado por (tipo, casas decimais)."""
    num = '0' + ('.' + '0' * dec if dec else '')
    return {'Número': '#,##' + num, 'Moeda': '"R$" #,##' + num, 'Porcentagem': num + '%',
            'Data': DATE_CODE}.get(nf, 'General')


def parse_nf(code):
    """Código do Excel -> (tipo, casas) ou None se não reconhecido."""
    if code in (None, '', 'General'):
        return 'Geral', 2
    s = re.sub(r'\[\$([^\]-]*)(?:-[^\]]*)?\]', r'\1', code.split(';')[0])  # [$R$-416] -> R$
    s = re.sub(r'\[[^\]]*\]', '', s)                                       # [Red], [$-416] ...
    s = re.sub(r'_.|\*.', '', s.replace('\\', '').replace('"', '')).strip()  # espaçamento do formato contábil
    if re.search(r'[dy]', s, re.I) and not re.search(r'[0#?]', s):
        return 'Data', 2
    m = re.fullmatch(r'(R\$)?\s*(#,##)?0(?:[.,](0{1,4}))?\s*(%)?', s)
    if not m or (m.group(1) and m.group(4)):
        return None
    kind = 'Moeda' if m.group(1) else 'Porcentagem' if m.group(4) else 'Número'
    return kind, len(m.group(3) or '')


EPOCH = datetime(1899, 12, 30)
DATE_RE = re.compile(r'(\d{1,2})/(\d{1,2})/(\d{4})(?:\s+(\d{1,2}):(\d{2})(?::(\d{2}))?)?')


class Date(float):
    """Data como número de série do Excel (conta como número nas fórmulas)."""


def to_serial(dt):
    return Date((dt - EPOCH) / timedelta(days=1))


def from_serial(v):
    return EPOCH + timedelta(seconds=round(v * 86400))


def parse_date(s):
    """'26/01/1991' ou '26/01/1991 08:30' -> Date; senão None."""
    m = DATE_RE.fullmatch(s.strip())
    if not m:
        return None
    d, mo, y, h, mi, sec = (int(g or 0) for g in m.groups())
    try:
        return to_serial(datetime(y, mo, d, h, mi, sec))
    except ValueError:
        return None


def fmt_date(v):
    dt = from_serial(v)
    s = dt.strftime('%d/%m/%Y')
    if dt.second:
        return s + dt.strftime(' %H:%M:%S')
    return s + dt.strftime(' %H:%M') if (dt.hour or dt.minute) else s


def col_name(c):
    s, c = '', c + 1
    while c:
        c, m = divmod(c - 1, 26)
        s = chr(65 + m) + s
    return s


def col_index(s):
    n = 0
    for ch in s.upper():
        n = n * 26 + ord(ch) - 64
    return n - 1


def cell_name(r, c):
    return f'{col_name(c)}{r + 1}'


class Err(str):
    """Valor de erro de célula (#ERRO!, #DIV/0!, #CIRC!)."""


class FErr(Exception):
    pass


def ref_rc(ref):
    m = re.fullmatch(r'\$?([A-Za-z]{1,3})\$?(\d+)', ref)
    if not m or int(m.group(2)) < 1:
        raise FErr('#ERRO!')
    return int(m.group(2)) - 1, col_index(m.group(1))


NUM_RE = re.compile(r'[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?')


def to_number(s):
    """'1.234,56' / '1,5' / '1.5' -> float; texto -> None."""
    s = s.strip()
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    return float(s) if NUM_RE.fullmatch(s) else None


def br(x, dec=2):
    return f'{x:,.{dec}f}'.translate(str.maketrans(',.', '.,'))


def fmt_number(v, nf='Geral', dec=2):
    if nf == 'Número':
        return br(v, dec)
    if nf == 'Moeda':
        return ('-R$ ' if v < 0 else 'R$ ') + br(abs(v), dec)
    if nf == 'Porcentagem':
        return br(v * 100, dec) + '%'
    if nf == 'Data' or isinstance(v, Date):
        try:
            return fmt_date(v)
        except (OverflowError, ValueError):
            pass
    if abs(v) < 1e15 and v == int(v):
        return str(int(v))
    return f'{v:.10g}'.replace('.', ',')


# --- fórmulas: tokenizador + descida recursiva (sem eval)
TOK = re.compile(r'\s*(?:(\d+\.?\d*(?:[eE][+-]?\d+)?|\.\d+)|(\$?[A-Za-z]{1,3}\$?\d+)(?![\w(])|([^\W\d]\w*)|(\S))')


def tokenize(src):
    out, pos, src = [], 0, src.rstrip()
    while pos < len(src):
        m = TOK.match(src, pos)
        if not m:
            raise FErr('#ERRO!')
        pos = m.end()
        num, ref, name, op = m.groups()
        if num:
            out.append(('num', float(num)))
        elif ref:
            out.append(('ref', ref))
        elif name:
            out.append(('name', name.upper()))
        else:
            out.append(('op', op))
    return out


class _Parser:
    def __init__(s, toks):
        s.t, s.i = toks, 0

    def peek(s):
        return s.t[s.i] if s.i < len(s.t) else (None, None)

    def take(s, op=None):
        tok = s.peek()
        if tok[0] is None or (op and tok != ('op', op)):
            raise FErr('#ERRO!')
        s.i += 1
        return tok

    def expr(s):
        n = s.term()
        while s.peek() in (('op', '+'), ('op', '-')):
            n = ('bin', s.take()[1], n, s.term())
        return n

    def term(s):
        n = s.unary()
        while s.peek() in (('op', '*'), ('op', '/')):
            n = ('bin', s.take()[1], n, s.unary())
        return n

    def unary(s):
        if s.peek() in (('op', '-'), ('op', '+')):
            op = s.take()[1]
            n = s.unary()
            return ('neg', n) if op == '-' else n
        return s.primary()

    def primary(s):
        kind, v = s.take()
        if kind == 'num':
            return ('num', v)
        if kind == 'ref':
            a = ref_rc(v)
            if s.peek() == ('op', ':'):
                s.take()
                k2, v2 = s.take()
                if k2 != 'ref':
                    raise FErr('#ERRO!')
                b = ref_rc(v2)
                return ('range', min(a[0], b[0]), min(a[1], b[1]), max(a[0], b[0]), max(a[1], b[1]))
            return ('ref',) + a
        if kind == 'name':
            if v not in FUNCS:
                raise FErr('#ERRO!')
            s.take('(')
            args = []
            if s.peek() != ('op', ')'):
                args.append(s.expr())
                while s.peek() in (('op', ','), ('op', ';')):
                    s.take()
                    args.append(s.expr())
            s.take(')')
            return ('fn', FUNCS[v], args)
        if (kind, v) == ('op', '('):
            n = s.expr()
            s.take(')')
            return n
        raise FErr('#ERRO!')


@lru_cache(maxsize=4096)
def parse(src):
    p = _Parser(tokenize(src))
    n = p.expr()
    if p.i != len(p.t):
        raise FErr('#ERRO!')
    return n


def evaluate(n, get):
    k = n[0]
    if k == 'num':
        return n[1]
    if k == 'ref':
        v = get(n[1], n[2])
        if isinstance(v, Err):
            raise FErr(v)
        if v is None:
            return 0.0
        if isinstance(v, float):
            return v
        raise FErr('#ERRO!')  # texto em conta
    if k == 'neg':
        return -evaluate(n[1], get)
    if k == 'bin':
        a, b = evaluate(n[2], get), evaluate(n[3], get)
        op = n[1]
        if op == '+':
            return a + b
        if op == '-':
            return a - b
        if op == '*':
            return a * b
        if b == 0:
            raise FErr('#DIV/0!')
        return a / b
    if k == 'range':
        raise FErr('#ERRO!')  # intervalo fora de função
    vals = []
    for a in n[2]:
        if a[0] == 'ref':
            a = ('range', a[1], a[2], a[1], a[2])
        if a[0] == 'range':
            for r in range(a[1], a[3] + 1):
                for c in range(a[2], a[4] + 1):
                    v = get(r, c)
                    if isinstance(v, Err):
                        raise FErr(v)
                    if isinstance(v, float):  # vazias e texto são ignoradas
                        vals.append(v)
        else:
            vals.append(evaluate(a, get))
    fn = n[1]
    if fn == 'SUM':
        return float(sum(vals))
    if fn == 'AVERAGE':
        if not vals:
            raise FErr('#DIV/0!')
        return sum(vals) / len(vals)
    if not vals:
        return 0.0
    return min(vals) if fn == 'MIN' else max(vals)


def to_excel_formula(raw):
    """'=soma(A1;B1)' -> '=SUM(A1,B1)'."""
    body = re.sub(r'([^\W\d]\w*)\s*\(', lambda m: FUNCS.get(m.group(1).upper(), m.group(1)) + '(', raw[1:])
    return '=' + body.replace(';', ',')


REF_RE = re.compile(r'(?<![\w$])(\$?)([A-Za-z]{1,3})(\$?)(\d+)(?![\w(])')


def shift_refs(formula, dr, dc):
    """Ajusta referências relativas (sem $) ao colar."""
    def rep(m):
        cabs, col, rabs, row = m.groups()
        c = col_index(col) + (0 if cabs else dc)
        r = int(row) - 1 + (0 if rabs else dr)
        if r < 0 or c < 0:
            return '#REF!'
        return f'{cabs}{col_name(c)}{rabs}{r + 1}'
    return REF_RE.sub(rep, formula)


def parse_tsv(text):
    lines = text.replace('\r\n', '\n').replace('\r', '\n').split('\n')
    if lines and lines[-1] == '':
        lines.pop()
    return [line.split('\t') for line in lines]


class Sheet:
    def __init__(self, name, nrows=100, ncols=26):
        self.name, self.nrows, self.ncols = name, nrows, ncols
        self.cells = {}   # (r, c) -> texto bruto digitado
        self.fmt = {}     # (r, c) -> Fmt
        self.colw, self.rowh = {}, {}  # coluna/linha -> pixels (só as diferentes do padrão)
        self.freeze = (0, 0)           # linhas, colunas congeladas
        self.src = None                # título da aba no arquivo de origem (None = aba nova)
        self.cache = {}
        self.busy = set()
        self.snapshot()

    def snapshot(self):
        """Guarda o estado igual ao do arquivo, para salvar só o que mudou."""
        self.base = (dict(self.cells), dict(self.fmt), dict(self.colw), dict(self.rowh), self.freeze)

    def raw(self, r, c):
        return self.cells.get((r, c), '')

    def getfmt(self, r, c):
        return self.fmt.get((r, c), DEF)

    def value(self, r, c):
        """None (vazia), float, str (texto) ou Err."""
        key = (r, c)
        if key in self.cache:
            return self.cache[key]
        raw = self.cells.get(key, '')
        if raw.startswith('='):
            if key in self.busy:
                raise FErr('#CIRC!')
            self.busy.add(key)
            try:
                v = evaluate(parse(raw[1:]), self.value)
            except FErr as e:
                v = Err(e.args[0])
            except RecursionError:
                v = Err('#ERRO!')
            finally:
                self.busy.discard(key)
        elif raw == '':
            v = None
        elif raw.startswith("'"):
            v = raw[1:]
        else:
            n = to_number(raw)
            if n is None:
                n = parse_date(raw)
            v = raw if n is None else n
        self.cache[key] = v
        return v

    def text(self, r, c):
        v = self.value(r, c)
        if v is None:
            return ''
        if isinstance(v, float):
            f = self.getfmt(r, c)
            return fmt_number(v, f.nf, f.dec)
        return str(v)

    def copy_block(self, r0, c0, r1, c1):
        return [[(self.raw(r, c), self.getfmt(r, c)) for c in range(c0, c1 + 1)] for r in range(r0, r1 + 1)]

    def to_tsv(self, r0, c0, r1, c1):
        return '\n'.join('\t'.join(self.text(r, c) for c in range(c0, c1 + 1)) for r in range(r0, r1 + 1))


BAD_NAME = re.compile(r'[\[\]:*?/\\]')


class Workbook:
    def __init__(self):
        self.sheets = [Sheet('Planilha1')]
        self.undo_stack, self.redo_stack = [], []
        self.dirty = False
        self.path = None
        self.src_path = None  # .xlsx usado como base ao salvar (preserva o que não exibimos)

    # --- planilhas (abas)
    def check_name(self, name, skip=None):
        name = name.strip()
        if not name or len(name) > 31 or BAD_NAME.search(name):
            raise ValueError('Nome inválido. Use até 31 caracteres, sem [ ] : * ? / \\')
        if any(s.name.lower() == name.lower() for s in self.sheets if s is not skip):
            raise ValueError(f'Já existe uma planilha chamada "{name}".')
        return name

    def add_sheet(self, name=None):
        if name is None:
            i = 1
            while any(s.name == f'Planilha{i}' for s in self.sheets):
                i += 1
            name = f'Planilha{i}'
        sh = Sheet(self.check_name(name))
        self.sheets.append(sh)
        self.dirty = True
        return sh

    def rename_sheet(self, i, name):
        self.sheets[i].name = self.check_name(name, skip=self.sheets[i])
        self.dirty = True

    def delete_sheet(self, i):
        if len(self.sheets) == 1:
            raise ValueError('A pasta de trabalho precisa ter pelo menos uma planilha.')
        del self.sheets[i]
        self.dirty = True

    # --- alterações (todas passam por change -> desfazer/refazer)
    @staticmethod
    def _apply(sh, state):
        for (r, c), (raw, fmt) in state.items():
            if raw:
                sh.cells[(r, c)] = raw
            else:
                sh.cells.pop((r, c), None)
            if fmt != DEF:
                sh.fmt[(r, c)] = fmt
            else:
                sh.fmt.pop((r, c), None)
            sh.nrows, sh.ncols = max(sh.nrows, r + 1), max(sh.ncols, c + 1)
        sh.cache.clear()  # ponytail: recalcula tudo a cada mudança; grafo de dependências se ficar lento

    def change(self, sh, updates):
        """updates: {(r, c): (bruto ou None, formato ou None)}; None = manter."""
        old = {k: (sh.raw(*k), sh.getfmt(*k)) for k in updates}
        new = {k: (old[k][0] if raw is None else raw, old[k][1] if fmt is None else fmt)
               for k, (raw, fmt) in updates.items()}
        if old == new:
            return
        self._apply(sh, new)
        self.undo_stack.append((sh, old, new))
        self.redo_stack.clear()
        self.dirty = True

    def undo(self):
        if not self.undo_stack:
            return None
        sh, old, new = self.undo_stack.pop()
        self._apply(sh, old)
        self.redo_stack.append((sh, old, new))
        self.dirty = True
        return sh

    def redo(self):
        if not self.redo_stack:
            return None
        sh, old, new = self.redo_stack.pop()
        self._apply(sh, new)
        self.undo_stack.append((sh, old, new))
        self.dirty = True
        return sh

    def set_cell(self, sh, r, c, raw):
        self.change(sh, {(r, c): (raw, None)})

    def clear(self, sh, r0, c0, r1, c1, formats=False):
        self.change(sh, {(r, c): ('', DEF if formats else None)
                         for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)})

    def set_format(self, sh, r0, c0, r1, c1, dec_delta=0, **kw):
        """kw: campos de Fmt a trocar (ex.: bold=True, color='FF0000', color=None = sem cor)."""
        up = {}
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                old = sh.getfmt(r, c)
                f = old._replace(**kw)
                if dec_delta and f.nf != 'Data':
                    f = f._replace(nf='Número' if f.nf == 'Geral' else f.nf, dec=max(0, min(4, f.dec + dec_delta)))
                if (f.nf, f.dec) != (old.nf, old.dec):
                    f = f._replace(code=None)  # o usuário mudou o formato: não vale mais o código original
                up[(r, c)] = (None, f)
        self.change(sh, up)

    def paste_block(self, sh, r, c, block, origin=None, shift=True):
        """Cola bloco [[(bruto, formato|None)]] em (r, c). Retorna a última célula colada."""
        dr, dc = (r - origin[0], c - origin[1]) if origin else (0, 0)
        up = {}
        for i, row in enumerate(block):
            for j, (raw, fmt) in enumerate(row):
                if shift and raw.startswith('=') and (dr or dc):
                    raw = shift_refs(raw, dr, dc)
                up[(r + i, c + j)] = (raw, fmt)
        self.change(sh, up)
        return r + len(block) - 1, c + max(len(row) for row in block) - 1


# --- xlsx
def num_to_raw(v):
    if isinstance(v, int) or (abs(v) < 1e15 and v == int(v)):
        return str(int(v))
    return repr(v).replace('.', ',')


def formula_supported(f):
    try:
        parse(f[1:])
        return True
    except FErr:
        return False


def xl_value(raw):
    """Texto bruto da célula -> valor para o openpyxl."""
    if raw.startswith('='):
        return to_excel_formula(raw)
    if raw.startswith("'"):
        return raw[1:]  # "'=PROCV(...)" importada volta como fórmula para o Excel
    if not raw:
        return None
    n = to_number(raw)
    if n is not None:
        return int(n) if abs(n) < 1e15 and n == int(n) else n
    d = parse_date(raw)
    return raw if d is None else from_serial(d)


def code_for(f, v):
    """Código de formato a gravar: o original, se o usuário não mudou o formato."""
    if f.code:
        return f.code
    if f.nf == 'Geral' and isinstance(v, datetime):
        return DATE_CODE + (' hh:mm' if v.time() != time() else '')
    return nf_code(f.nf, f.dec)


def px_to_width(px):
    return round((px - 5) / 7, 2)


def set_col_width(ws, c, px):
    """Grava a largura da coluna c (1 = A), separando faixas <col min..max> do arquivo que a contenham."""
    from openpyxl.utils import get_column_letter, column_index_from_string
    for d in list(ws.column_dimensions.values()):
        lo = d.min or column_index_from_string(d.index)
        hi = d.max or lo
        if lo < hi and lo <= c <= hi:
            d.min = d.max = lo
            for k in range(lo + 1, hi + 1):
                nd = copy(d)
                nd.index, nd.min, nd.max = get_column_letter(k), k, k
                ws.column_dimensions[nd.index] = nd
    ws.column_dimensions[get_column_letter(c)].width = px_to_width(px)


def write_sheet(sh, ws, base):
    """Grava na aba do openpyxl só o que difere de base (estado do arquivo de origem)."""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import PatternFill
    from openpyxl.styles.colors import Color
    cells0, fmt0, colw0, rowh0, freeze0 = base
    for (r, c) in sorted(set(sh.cells) | set(sh.fmt) | set(cells0) | set(fmt0)):
        raw, f = sh.raw(r, c), sh.getfmt(r, c)
        o = fmt0.get((r, c), DEF)
        changed = raw != cells0.get((r, c), '')
        if not changed and f == o:
            continue
        cell = ws.cell(row=r + 1, column=c + 1)
        if isinstance(cell, MergedCell):
            continue  # ponytail: só a 1ª célula de uma mesclagem é gravável
        v = xl_value(raw)
        if changed:
            cell.value = v
        if (f.bold, f.italic, f.color) != (o.bold, o.italic, o.color):
            ft = copy(cell.font)
            ft.b, ft.i = f.bold, f.italic
            if f.color != o.color:
                ft.color = Color(rgb='FF' + f.color) if f.color else None
            cell.font = ft
        if f.fill != o.fill:
            cell.fill = PatternFill('solid', fgColor='FF' + f.fill) if f.fill else PatternFill()
        if (f.align, f.wrap, f.valign) != (o.align, o.wrap, o.valign):
            al = copy(cell.alignment)
            if f.align != o.align:
                al.horizontal = f.align
            if f.wrap != o.wrap:
                al.wrap_text = f.wrap or None
            if f.valign != o.valign:
                al.vertical = f.valign
            cell.alignment = al
        if changed or (f.nf, f.dec, f.code) != (o.nf, o.dec, o.code):
            cell.number_format = code_for(f, v)
    for c in sh.colw:
        if sh.colw[c] != colw0.get(c):
            set_col_width(ws, c + 1, sh.colw[c])
    for r in sh.rowh:
        if sh.rowh[r] != rowh0.get(r):
            ws.row_dimensions[r + 1].height = round(sh.rowh[r] * 0.75, 2)
    if sh.freeze != freeze0:
        ws.freeze_panes = cell_name(*sh.freeze) if sh.freeze != (0, 0) else None


def save_xlsx(wb, path):
    """Salva por cima do .xlsx de origem (o que não exibimos fica preservado).
    Sem origem legível, grava do zero. Retorna True se usou a origem."""
    import openpyxl
    x = None
    if wb.src_path:
        try:
            with warnings.catch_warnings():
                warnings.simplefilter('ignore')
                x = openpyxl.load_workbook(wb.src_path)
        except Exception:
            x = None
    based = x is not None
    if not based:
        x = openpyxl.Workbook()
        x.remove(x.active)
    empty = ({}, {}, {}, {}, (0, 0))
    srcs = {sh.src for sh in wb.sheets} if based else set()
    for ws in list(x.worksheets):
        if ws.title not in srcs:
            x.remove(ws)                      # aba excluída aqui
    kept = {ws.title: ws for ws in x.worksheets}
    for i, ws in enumerate(kept.values()):
        ws.title = f'__tmp{i}__'              # renomeia em 2 passos (permite trocar nomes A<->B)
    out = []
    for sh in wb.sheets:
        ws = kept.get(sh.src) if based else None
        if ws is None:
            ws = x.create_sheet(sh.name)
        ws.title = sh.name
        out.append((sh, ws))
    x._sheets = [ws for _, ws in out] + list(x.chartsheets)  # ordem das abas como aqui
    x.active = 0
    for i, (sh, ws) in enumerate(out):
        ws.sheet_view.tabSelected = i == 0
        write_sheet(sh, ws, sh.base if (based and sh.src) else empty)
    x.save(path)
    wb.src_path = path                         # o arquivo salvo vira a nova base
    for sh in wb.sheets:
        sh.src = sh.name
        sh.snapshot()
    return based


def xl_color(c):
    """Color do openpyxl -> 'RRGGBB' ou None (cores de tema não são exibidas)."""
    from openpyxl.styles.colors import COLOR_INDEX
    if c is None:
        return None
    if c.type == 'rgb' and isinstance(c.rgb, str):
        return c.rgb[-6:].upper()
    if c.type == 'indexed' and isinstance(c.indexed, int) and c.indexed < len(COLOR_INDEX) - 2:
        return COLOR_INDEX[c.indexed][-6:].upper()  # 64/65 = cores do sistema
    return None


LOST = ' — será perdido ao salvar'


def load_xlsx(path):
    """Retorna (Workbook, relatório). No relatório, itens que terminam com LOST se perdem ao salvar;
    os demais ficam preservados no arquivo, mas não aparecem aqui."""
    import openpyxl
    with warnings.catch_warnings(record=True) as warns:
        warnings.simplefilter('always')
        x = openpyxl.load_workbook(path)
    wb = Workbook()
    wb.sheets = []
    rep, bad_funcs, bad_formulas = set(), set(), 0
    for w in warns:
        if 'not supported' in str(w.message) or 'will be removed' in str(w.message):
            rep.add(f'Recurso não suportado pelo leitor ({w.message})' + LOST)
    if len(x.defined_names):
        rep.add('Nomes definidos')
    if x.chartsheets:
        rep.add('Planilhas de gráfico' + LOST)
    for ws in x.worksheets:
        sh = Sheet(ws.title, max(100, ws.max_row), max(26, ws.max_column))
        wb.sheets.append(sh)
        checks = [(ws.merged_cells.ranges, 'Células mescladas (só a 1ª célula de cada mesclagem pode ser editada)'),
                  (ws._charts, 'Gráficos' + LOST), (ws._images, 'Imagens' + LOST),
                  (list(ws.conditional_formatting), 'Formatação condicional'),
                  (ws.data_validations.dataValidation, 'Validação de dados'), (ws.auto_filter.ref, 'Filtros'),
                  (ws.tables, 'Tabelas'), (ws.protection.sheet, 'Proteção de planilha')]
        rep.update(label for found, label in checks if found)
        for d in ws.column_dimensions.values():
            if d.customWidth and d.min:
                for c in range(d.min, (d.max or d.min) + 1):
                    sh.colw[c - 1] = round(d.width * 7 + 5)
        for r, d in ws.row_dimensions.items():
            if d.customHeight and d.ht:
                sh.rowh[r - 1] = round(d.ht * 4 / 3)
        if ws.freeze_panes:
            sh.freeze = ref_rc(ws.freeze_panes)
        for row in ws.iter_rows():
            for cell in row:
                if type(cell).__name__ == 'MergedCell':
                    continue
                font, al, fill = cell.font, cell.alignment, cell.fill
                fillc = None
                if fill is not None and fill.fill_type:
                    fillc = xl_color(fill.fgColor) if fill.fill_type == 'solid' else None
                    if fillc is None:
                        rep.add('Preenchimentos com padrão ou cor de tema (exibidos sem cor)')
                color = xl_color(font.color)
                if font.color is not None and color is None and font.color.type == 'theme' and font.color.theme != 1:
                    rep.add('Cores de fonte do tema (exibidas em preto)')
                if color == '000000':
                    color = None
                if font.u or font.strike:
                    rep.add('Sublinhado / tachado')
                if any(getattr(cell.border, side).style for side in ('left', 'right', 'top', 'bottom')):
                    rep.add('Bordas')
                if al.horizontal not in (None, 'general', 'left', 'center', 'right') or al.textRotation or al.indent:
                    rep.add('Alinhamento justificado, recuo ou texto girado')
                if cell.comment:
                    rep.add('Comentários')
                if cell.hyperlink:
                    rep.add('Hiperlinks')
                parsed = parse_nf(cell.number_format)
                if parsed is None:
                    if cell.value is not None:
                        rep.add(f'Formato de número "{cell.number_format}" (exibido como Geral)')
                    parsed = ('Geral', 2)
                code = None if nf_code(*parsed) == cell.number_format else cell.number_format
                v = cell.value
                f = Fmt(bool(font.b), bool(font.i), parsed[0], parsed[1], color, fillc,
                        al.horizontal if al.horizontal in ('left', 'center', 'right') else None,
                        bool(al.wrap_text), al.vertical if al.vertical in ('top', 'center', 'bottom') else None, code)
                r, c = cell.row - 1, cell.column - 1
                if f != DEF:
                    sh.fmt[(r, c)] = f
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith('='):
                    if formula_supported(v):
                        raw = v
                    else:
                        raw = "'" + v
                        bad_formulas += 1
                        bad_funcs.update(n.upper() for n in re.findall(r'([^\W\d][\w.]*)\s*\(', v)
                                         if n.upper() not in FUNCS)
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    raw = num_to_raw(v)
                elif isinstance(v, date):  # datetime também é date
                    raw = fmt_date(to_serial(v if isinstance(v, datetime) else datetime(v.year, v.month, v.day)))
                elif isinstance(v, str):
                    raw = "'" + v if (to_number(v) is not None or parse_date(v) is not None
                                      or v.startswith(("'", '='))) else v
                else:
                    rep.add('Horas, booleanos ou fórmulas matriciais (exibidos como texto; preservados se não editados)')
                    raw = "'" + str(getattr(v, 'text', v))
                sh.cells[(r, c)] = raw
        sh.src = ws.title
        sh.snapshot()
    if bad_formulas:
        funcs = f" (funções: {', '.join(sorted(bad_funcs))})" if bad_funcs else ''
        rep.add(f'{bad_formulas} fórmula(s) não suportada(s){funcs} — exibidas como texto; '
                'continuam como fórmula no arquivo')
    if not wb.sheets:
        wb.sheets = [Sheet('Planilha1')]
    wb.path = wb.src_path = path
    return wb, sorted(rep)


# ---------------------------------------------------------------- interface
CW, RH, HW, HH = 90, 22, 50, 22  # largura/altura da célula, cabeçalhos


class App:
    def __init__(self, root, wb=None):
        self.root = root
        self.wb = wb or Workbook()
        self.si = 0
        self.cur = self.end = (0, 0)
        self.top = self.left = 0
        self.clip = None
        self.editor = None
        self.resizing = None
        base = tkfont.nametofont('TkDefaultFont').actual()
        self.fonts = {(b, i): tkfont.Font(root=root, family=base['family'], size=10,
                                          weight='bold' if b else 'normal', slant='italic' if i else 'roman')
                      for b in (False, True) for i in (False, True)}
        self._build()
        self.build_tabs()
        self.refresh()

    @property
    def sh(self):
        return self.wb.sheets[self.si]

    def rect(self):
        (r0, c0), (r1, c1) = self.cur, self.end
        return min(r0, r1), min(c0, c1), max(r0, r1), max(c0, c1)

    # --- montagem
    def _build(self):
        r = self.root
        r.geometry('1000x650')
        menu = tk.Menu(r)
        r.config(menu=menu)
        def submenu(label, items):
            m = tk.Menu(menu, tearoff=0)
            menu.add_cascade(label=label, menu=m)
            for it in items:
                if it is None:
                    m.add_separator()
                else:
                    m.add_command(label=it[0], command=it[1], accelerator=it[2] if len(it) > 2 else '')
        submenu('Arquivo', [('Novo', self.new, 'Ctrl+N'), ('Abrir...', self.open, 'Ctrl+O'),
                            ('Salvar', self.save, 'Ctrl+S'), ('Salvar como...', self.save_as), None,
                            ('Sair', self.on_close)])
        submenu('Editar', [('Desfazer', self.undo, 'Ctrl+Z'), ('Refazer', self.redo, 'Ctrl+Y'), None,
                           ('Recortar', self.cut, 'Ctrl+X'), ('Copiar', self.copy, 'Ctrl+C'),
                           ('Colar', self.paste, 'Ctrl+V'), ('Limpar', self.clear, 'Del')])
        submenu('Planilha', [('Nova planilha', self.add_sheet), ('Renomear planilha...', lambda: self.rename_sheet(self.si)),
                             ('Excluir planilha', self.delete_sheet)])
        submenu('Exibir', [('Congelar painéis (acima e à esquerda da célula atual)', self.freeze_panes),
                           ('Descongelar painéis', self.unfreeze)])

        tb = tk.Frame(r, pady=2)
        tb.pack(fill='x')
        tk.Button(tb, text='N', width=3, font=self.fonts[(True, False)], command=self.toggle_bold).pack(side='left', padx=2)
        tk.Button(tb, text='I', width=3, font=self.fonts[(False, True)], command=self.toggle_italic).pack(side='left', padx=2)
        tk.Label(tb, text='  Formato:').pack(side='left')
        self.nfvar = tk.StringVar(value='Geral')
        cb = ttk.Combobox(tb, textvariable=self.nfvar, values=FORMATS, state='readonly', width=14)
        cb.pack(side='left')
        cb.bind('<<ComboboxSelected>>', lambda e: (self.apply_fmt(nf=self.nfvar.get()), self.cv.focus_set()))
        def btn(text, cmd, padx=2):
            tk.Button(tb, text=text, command=cmd).pack(side='left', padx=padx)
        btn('− decimais', lambda: self.apply_fmt(dec_delta=-1), (6, 2))
        btn('+ decimais', lambda: self.apply_fmt(dec_delta=1))
        btn('Esquerda', lambda: self.apply_fmt(align='left'), (10, 0))
        btn('Centro', lambda: self.apply_fmt(align='center'), 0)
        btn('Direita', lambda: self.apply_fmt(align='right'), 0)
        btn('Quebrar texto', lambda: self.apply_fmt(wrap=not self.sh.getfmt(*self.cur).wrap), (6, 2))
        for label, field in (('Cor do texto', 'color'), ('Cor de fundo', 'fill')):
            mb = tk.Menubutton(tb, text=label + ' ▾', relief='raised', padx=6, pady=3)
            m = tk.Menu(mb, tearoff=0)
            m.add_command(label='Escolher cor...', command=lambda f=field, t=label: self.pick_color(f, t))
            m.add_command(label='Sem cor', command=lambda f=field: self.apply_fmt(**{f: None}))
            mb.config(menu=m)
            mb.pack(side='left', padx=(6, 0))

        fb = tk.Frame(r)
        fb.pack(fill='x')
        self.namelbl = tk.Label(fb, width=10, relief='sunken', anchor='w')
        self.namelbl.pack(side='left', padx=2)
        tk.Label(fb, text='fx', font=self.fonts[(False, True)]).pack(side='left')
        self.fbar = tk.Entry(fb, font=self.fonts[(False, False)])
        self.fbar.pack(side='left', fill='x', expand=True, padx=2)
        self.fbar.bind('<Return>', self.commit_fbar)
        self.fbar.bind('<Escape>', lambda e: (self.refresh(), self.cv.focus_set()))

        self.tabbar = tk.Frame(r)
        self.tabbar.pack(side='bottom', fill='x')
        gf = tk.Frame(r)
        gf.pack(fill='both', expand=True)
        self.vs = tk.Scrollbar(gf, orient='vertical', command=self.yview)
        self.vs.pack(side='right', fill='y')
        self.hs = tk.Scrollbar(gf, orient='horizontal', command=self.xview)
        self.hs.pack(side='bottom', fill='x')
        self.cv = cv = tk.Canvas(gf, bg='white', highlightthickness=0, takefocus=1)
        cv.pack(fill='both', expand=True)

        cv.bind('<Configure>', lambda e: self.redraw())
        cv.bind('<Button-1>', self.on_click)
        cv.bind('<Shift-Button-1>', lambda e: self.on_drag(e))
        cv.bind('<B1-Motion>', self.on_drag)
        cv.bind('<ButtonRelease-1>', self.on_release)
        cv.bind('<Motion>', lambda e: cv.config(cursor={'c': 'sb_h_double_arrow', 'r': 'sb_v_double_arrow'}.get(
            (self.border_at(e) or ' ')[0], '')))
        cv.bind('<Double-Button-1>', self.on_double)
        cv.bind('<MouseWheel>', lambda e: self.scroll_rows(-3 if e.delta > 0 else 3))
        cv.bind('<Key>', self.on_key)
        for key, fn in (('c', self.copy), ('x', self.cut), ('v', self.paste), ('z', self.undo),
                        ('y', self.redo), ('b', self.toggle_bold), ('i', self.toggle_italic), ('a', self.select_all)):
            for k in (key, key.upper()):
                cv.bind(f'<Control-{k}>', lambda e, fn=fn: (fn(), 'break')[1])
        for k, fn in (('n', self.new), ('o', self.open), ('s', self.save)):
            for kk in (k, k.upper()):
                r.bind(f'<Control-{kk}>', lambda e, fn=fn: (fn(), 'break')[1])
        r.protocol('WM_DELETE_WINDOW', self.on_close)
        cv.focus_set()

    # --- desenho
    def axis(self, cols):
        """[(índice, posição, tamanho)] das colunas (ou linhas) visíveis: congeladas + a partir da rolagem."""
        sh = self.sh
        if cols:
            n, frozen, start, size, dflt, head, limit = (sh.ncols, sh.freeze[1], self.left, sh.colw, CW, HW,
                                                         max(self.cv.winfo_width(), 300))
        else:
            n, frozen, start, size, dflt, head, limit = (sh.nrows, sh.freeze[0], self.top, sh.rowh, RH, HH,
                                                         max(self.cv.winfo_height(), 200))
        out, pos = [], head
        for i in list(range(min(frozen, n))) + list(range(max(start, frozen), n)):
            if pos >= limit:
                break
            s = size.get(i, dflt)
            out.append((i, pos, s))
            pos += s
        return out

    def visible(self):
        """Quantas linhas/colunas roláveis cabem na tela."""
        fr, fc = self.sh.freeze
        return (max(1, sum(1 for i, _, _ in self.axis(False) if i >= fr) - 1),
                max(1, sum(1 for i, _, _ in self.axis(True) if i >= fc) - 1))

    def wrap_lines(self, txt, font, width):
        lines = []
        for para in txt.split('\n'):
            line = ''
            for word in para.split(' '):
                test = f'{line} {word}' if line else word
                if line and font.measure(test) > width:
                    lines.append(line)
                    line = word
                else:
                    line = test
            lines.append(line)
        return lines

    def redraw(self):
        cv, sh = self.cv, self.sh
        cv.delete('all')
        cols, rows = self.axis(True), self.axis(False)
        r0, c0, r1, c1 = self.rect()
        for c, x, w in cols:
            cv.create_rectangle(x, 0, x + w, HH, fill='#cfd8e8' if c0 <= c <= c1 else '#eeeeee', outline='#bbbbbb')
            cv.create_text(x + w / 2, HH / 2, text=col_name(c))
        for r, y, h in rows:
            cv.create_rectangle(0, y, HW, y + h, fill='#cfd8e8' if r0 <= r <= r1 else '#eeeeee', outline='#bbbbbb')
            cv.create_text(HW / 2, y + h / 2, text=str(r + 1))
            for c, x, w in cols:
                f = sh.getfmt(r, c)
                sel = r0 <= r <= r1 and c0 <= c <= c1
                bg = '#' + f.fill if f.fill else ('#dde8f8' if sel else 'white')
                cv.create_rectangle(x, y, x + w, y + h, fill=bg, outline='#e2e2e2')
                txt = sh.text(r, c)
                if not txt:
                    continue
                font = self.fonts[(f.bold, f.italic)]
                align = f.align or ('right' if isinstance(sh.value(r, c), float) else 'left')
                if f.wrap:
                    ls = font.metrics('linespace')
                    txt = '\n'.join(self.wrap_lines(txt, font, w - 8)[:max(1, (h - 2) // ls)])
                va = f.valign or 'center'
                ty = {'top': y + 2, 'bottom': y + h - 2}.get(va, y + h / 2)
                tx = {'left': x + 4, 'right': x + w - 4}.get(align, x + w / 2)
                anchor = ({'top': 'n', 'bottom': 's'}.get(va, '') + {'left': 'w', 'right': 'e'}.get(align, '')) or 'center'
                cv.create_text(tx, ty, text=txt, anchor=anchor, justify=align, font=font,
                               fill='#' + f.color if f.color else 'black')
        fr, fc = sh.freeze
        if fr and rows and fr <= len(rows):
            y = rows[fr - 1][1] + rows[fr - 1][2]
            cv.create_line(0, y, cols[-1][1] + cols[-1][2] if cols else HW, y, fill='#555555', width=2)
        if fc and cols and fc <= len(cols):
            x = cols[fc - 1][1] + cols[fc - 1][2]
            cv.create_line(x, 0, x, rows[-1][1] + rows[-1][2] if rows else HH, fill='#555555', width=2)

        def span(items, a, b):
            inside = [(p, s) for i, p, s in items if a <= i <= b]
            return (inside[0][0], inside[-1][0] + inside[-1][1]) if inside else None

        def box(ra, ca, rb, cb, width, color):
            xs, ys = span(cols, ca, cb), span(rows, ra, rb)
            if xs and ys:
                cv.create_rectangle(xs[0], ys[0], xs[1], ys[1], outline=color, width=width)
        box(r0, c0, r1, c1, 1, '#4472c4')
        box(*self.cur, *self.cur, 2, '#1f4e9a')
        vr, vc = self.visible()
        self.vs.set(self.top / sh.nrows, min(1, (self.top + vr) / sh.nrows))
        self.hs.set(self.left / sh.ncols, min(1, (self.left + vc) / sh.ncols))

    def refresh(self):
        r0, c0, r1, c1 = self.rect()
        self.namelbl.config(text=cell_name(r0, c0) + (f':{cell_name(r1, c1)}' if (r0, c0) != (r1, c1) else ''))
        try:
            focus = self.root.focus_get()
        except (KeyError, tk.TclError):  # foco em janela interna do Tk (ex.: lista do combobox)
            focus = None
        if focus is not self.fbar:
            self.fbar.delete(0, 'end')
            self.fbar.insert(0, self.sh.raw(*self.cur))
        self.nfvar.set(self.sh.getfmt(*self.cur)[2])
        name = os.path.basename(self.wb.path) if self.wb.path else 'Sem título'
        self.root.title(f"{name}{'*' if self.wb.dirty else ''} - planilha_code")
        self.redraw()

    def build_tabs(self):
        for w in self.tabbar.winfo_children():
            w.destroy()
        self.tabbtns = []
        for i, sh in enumerate(self.wb.sheets):
            b = tk.Button(self.tabbar, text=sh.name, padx=8, command=lambda i=i: self.switch(i))
            b.bind('<Double-Button-1>', lambda e, i=i: self.rename_sheet(i))
            b.pack(side='left')
            self.tabbtns.append(b)
        tk.Button(self.tabbar, text='+', width=3, command=self.add_sheet).pack(side='left', padx=4)
        self.style_tabs()

    def style_tabs(self):
        for i, b in enumerate(self.tabbtns):
            b.config(relief='sunken' if i == self.si else 'raised', bg='white' if i == self.si else '#dddddd',
                     font=self.fonts[(i == self.si, False)])

    # --- navegação e seleção
    def ensure_visible(self, rc):
        sh = self.sh
        self.top = self._scroll_to(rc[0], sh.freeze[0], self.top, sh.rowh, RH, max(self.cv.winfo_height(), 200) - HH)
        self.left = self._scroll_to(rc[1], sh.freeze[1], self.left, sh.colw, CW, max(self.cv.winfo_width(), 300) - HW)

    @staticmethod
    def _scroll_to(i, frozen, start, size, dflt, avail):
        """Novo início de rolagem para que o índice i fique inteiro na tela."""
        start = max(start, frozen)
        if i < frozen:
            return start
        if i < start:
            return i
        avail -= sum(size.get(k, dflt) for k in range(frozen))
        used, s = 0, i
        while s >= frozen:
            used += size.get(s, dflt)
            if used > avail:
                break
            s -= 1
        return max(start, min(i, s + 1))

    def move(self, dr, dc, extend=False):
        r, c = self.end if extend else self.cur
        r, c = max(0, r + dr), max(0, c + dc)
        self.sh.nrows, self.sh.ncols = max(self.sh.nrows, r + 1), max(self.sh.ncols, c + 1)
        if extend:
            self.end = (r, c)
        else:
            self.cur = self.end = (r, c)
        self.ensure_visible((r, c))
        self.refresh()

    @staticmethod
    def _find(items, p, head):
        if p < head:
            return -1
        for i, pos, s in items:
            if pos <= p < pos + s:
                return i
        return items[-1][0] if items else -1

    def hit(self, e):
        return self._find(self.axis(False), e.y, HH), self._find(self.axis(True), e.x, HW)

    def border_at(self, e):
        """('c', coluna, x inicial) / ('r', linha, y inicial) se o mouse está na borda de um cabeçalho."""
        if e.y < HH:
            for c, x, w in self.axis(True):
                if abs(e.x - (x + w)) <= 3:
                    return 'c', c, x
        elif e.x < HW:
            for r, y, h in self.axis(False):
                if abs(e.y - (y + h)) <= 3:
                    return 'r', r, y
        return None

    def on_click(self, e):
        self.finish_edit()
        self.cv.focus_set()
        self.resizing = self.border_at(e)
        if self.resizing:
            return
        r, c = self.hit(e)
        if r < 0 and c < 0:
            return self.select_all()
        if r < 0:
            self.cur, self.end = (0, c), (self.sh.nrows - 1, c)
        elif c < 0:
            self.cur, self.end = (r, 0), (r, self.sh.ncols - 1)
        else:
            self.cur = self.end = (r, c)
        self.refresh()

    def on_drag(self, e):
        if self.resizing:
            kind, i, start = self.resizing
            if kind == 'c':
                self.sh.colw[i] = max(8, e.x - start)
            else:
                self.sh.rowh[i] = max(8, e.y - start)
            self.wb.dirty = True
            return self.redraw()
        r, c = self.hit(e)
        if r >= 0 and c >= 0 and (r, c) != self.end:
            self.end = (r, c)
            self.refresh()

    def on_release(self, e):
        if self.resizing:
            self.resizing = None
            self.refresh()

    def on_double(self, e):
        b = self.border_at(e)
        if not b:
            return self.start_edit()
        self.resizing = None
        if b[0] == 'c':
            self.autofit(b[1])

    def autofit(self, c):
        sh = self.sh
        widths = [self.fonts[sh.getfmt(r, cc)[:2]].measure(sh.text(r, cc)) for (r, cc) in sh.cells if cc == c]
        sh.colw[c] = max(30, max(widths, default=0) + 12)
        self.wb.dirty = True
        self.refresh()

    def freeze_panes(self):
        self.sh.freeze = self.cur
        self.top, self.left = max(self.top, self.cur[0]), max(self.left, self.cur[1])
        self.wb.dirty = True
        self.refresh()

    def unfreeze(self):
        self.sh.freeze = (0, 0)
        self.wb.dirty = True
        self.refresh()

    def select_all(self):
        self.cur, self.end = (0, 0), (self.sh.nrows - 1, self.sh.ncols - 1)
        self.refresh()

    def scroll_rows(self, n):
        self.top = max(self.sh.freeze[0], min(self.sh.nrows - 1, self.top + n))
        self.redraw()

    def yview(self, *a):
        vr, _ = self.visible()
        if a[0] == 'moveto':
            self.top = int(float(a[1]) * self.sh.nrows)
        else:
            self.top += int(a[1]) * (vr if a[2] == 'pages' else 1)
        self.top = max(self.sh.freeze[0], min(self.sh.nrows - 1, self.top))
        self.redraw()

    def xview(self, *a):
        _, vc = self.visible()
        if a[0] == 'moveto':
            self.left = int(float(a[1]) * self.sh.ncols)
        else:
            self.left += int(a[1]) * (vc if a[2] == 'pages' else 1)
        self.left = max(self.sh.freeze[1], min(self.sh.ncols - 1, self.left))
        self.redraw()

    def on_key(self, e):
        k, shift = e.keysym, bool(e.state & 1)
        ctrl, alt = bool(e.state & 4), bool(e.state & 0x20000)
        moves = {'Up': (-1, 0), 'Down': (1, 0), 'Left': (0, -1), 'Right': (0, 1)}
        vr, _ = self.visible()
        if k in moves:
            self.move(*moves[k], extend=shift)
        elif k in ('Return', 'KP_Enter'):
            self.move(-1 if shift else 1, 0)
        elif k == 'Tab':
            self.move(0, -1 if shift else 1)
        elif k in ('Prior', 'Next'):
            self.move(-vr if k == 'Prior' else vr, 0, extend=shift)
        elif k == 'Home':
            self.move(0, -self.cur[1])
        elif k == 'F2':
            self.start_edit()
        elif k == 'Delete':
            self.clear()
        elif k == 'BackSpace':
            self.start_edit('')
        elif e.char and e.char.isprintable() and (not ctrl or alt):
            self.start_edit(e.char)
        else:
            return None
        return 'break'

    # --- edição
    def start_edit(self, text=None):
        self.finish_edit()
        self.end = self.cur
        self.ensure_visible(self.cur)
        self.refresh()
        r, c = self.cur
        x, w = next(((p, s) for i, p, s in self.axis(True) if i == c), (HW, CW))
        y, h = next(((p, s) for i, p, s in self.axis(False) if i == r), (HH, RH))
        e = tk.Entry(self.cv, relief='solid', bd=1, font=self.fonts[self.sh.getfmt(r, c)[:2]])
        e.insert(0, self.sh.raw(r, c) if text is None else text)
        e.place(x=x, y=y, width=max(w, 2 * CW if text is None else 0) + 1, height=min(h, 2 * RH) + 1)
        e.focus_set()
        e.icursor('end')
        for seq, d in (('<Return>', (1, 0)), ('<KP_Enter>', (1, 0)), ('<Tab>', (0, 1)), ('<Shift-Tab>', (0, -1)),
                       ('<Up>', (-1, 0)), ('<Down>', (1, 0))):
            e.bind(seq, lambda ev, d=d: self.finish_edit(*d))
        e.bind('<Escape>', lambda ev: self.cancel_edit())
        e.bind('<FocusOut>', lambda ev: self.finish_edit(refocus=False))
        e.bind('<KeyRelease>', lambda ev: (self.fbar.delete(0, 'end'), self.fbar.insert(0, e.get())))
        self.editor = (e, self.sh, r, c)

    def finish_edit(self, dr=0, dc=0, refocus=True):
        if not self.editor:
            return 'break'
        e, sh, r, c = self.editor
        self.editor = None
        txt = e.get()
        e.destroy()
        self.wb.change(sh, {(r, c): (txt, None)})
        if refocus:
            self.cv.focus_set()
        self.move(dr, dc) if (dr or dc) else self.refresh()
        return 'break'

    def cancel_edit(self):
        if self.editor:
            e = self.editor[0]
            self.editor = None
            e.destroy()
            self.cv.focus_set()
            self.refresh()
        return 'break'

    def commit_fbar(self, e=None):
        self.wb.set_cell(self.sh, *self.cur, self.fbar.get())
        self.cv.focus_set()
        self.move(1, 0)
        return 'break'

    # --- comandos de edição
    def undo(self):
        self.finish_edit()
        self._after_history(self.wb.undo())

    def redo(self):
        self.finish_edit()
        self._after_history(self.wb.redo())

    def _after_history(self, sh):
        if sh in self.wb.sheets and self.wb.sheets.index(sh) != self.si:
            self.switch(self.wb.sheets.index(sh))
        self.refresh()

    def clear(self):
        self.finish_edit()
        self.wb.clear(self.sh, *self.rect())
        self.refresh()

    def copy(self, cut=False):
        self.finish_edit()
        r0, c0, r1, c1 = self.rect()
        text = self.sh.to_tsv(r0, c0, r1, c1)
        self.clip = dict(block=self.sh.copy_block(r0, c0, r1, c1), origin=(r0, c0), text=text, cut=cut)
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        if cut:
            self.wb.clear(self.sh, r0, c0, r1, c1, formats=True)
            self.refresh()

    def cut(self):
        self.copy(cut=True)

    def paste(self):
        self.finish_edit()
        try:
            text = self.root.clipboard_get().replace('\r\n', '\n')
        except tk.TclError:
            text = None
        r, c = self.rect()[:2]
        clip = self.clip
        if clip and (text is None or text.rstrip('\n') == clip['text'].rstrip('\n')):
            end = self.wb.paste_block(self.sh, r, c, clip['block'], clip['origin'], shift=not clip['cut'])
        elif text:
            end = self.wb.paste_block(self.sh, r, c, [[(v, None) for v in row] for row in parse_tsv(text)])
        else:
            return
        self.cur, self.end = (r, c), end
        self.refresh()

    def apply_fmt(self, **kw):
        self.finish_edit()
        self.wb.set_format(self.sh, *self.rect(), **kw)
        self.refresh()

    def pick_color(self, field, title):
        cur = getattr(self.sh.getfmt(*self.cur), field)
        rgb = colorchooser.askcolor(color='#' + cur if cur else None, title=title, parent=self.root)[1]
        if rgb:
            self.apply_fmt(**{field: rgb[1:].upper()})

    def toggle_bold(self):
        self.apply_fmt(bold=not self.sh.getfmt(*self.cur)[0])

    def toggle_italic(self):
        self.apply_fmt(italic=not self.sh.getfmt(*self.cur)[1])

    # --- abas
    def switch(self, i):
        self.finish_edit(refocus=False)
        self.si = i
        self.cur = self.end = (0, 0)
        self.top = self.left = 0
        self.style_tabs()
        self.refresh()
        self.cv.focus_set()

    def add_sheet(self):
        self.finish_edit()
        self.wb.add_sheet()
        self.build_tabs()
        self.switch(len(self.wb.sheets) - 1)

    def rename_sheet(self, i):
        name = simpledialog.askstring('Renomear planilha', 'Novo nome:', initialvalue=self.wb.sheets[i].name,
                                      parent=self.root)
        if name is None:
            return
        try:
            self.wb.rename_sheet(i, name)
        except ValueError as ex:
            messagebox.showerror('Renomear planilha', str(ex), parent=self.root)
            return
        self.build_tabs()
        self.refresh()

    def delete_sheet(self):
        name = self.sh.name
        if not messagebox.askyesno('Excluir planilha', f'Excluir a planilha "{name}"?\nEsta ação não pode ser desfeita.',
                                   parent=self.root):
            return
        try:
            self.wb.delete_sheet(self.si)
        except ValueError as ex:
            messagebox.showerror('Excluir planilha', str(ex), parent=self.root)
            return
        self.si = min(self.si, len(self.wb.sheets) - 1)
        self.build_tabs()
        self.switch(self.si)

    # --- arquivo
    def confirm_discard(self):
        if not self.wb.dirty:
            return True
        ans = messagebox.askyesnocancel('planilha_code', 'Deseja salvar as alterações antes de continuar?', parent=self.root)
        if ans is None:
            return False
        return self.save() if ans else True

    def load_wb(self, wb):
        self.wb, self.si, self.clip = wb, 0, None
        self.cur = self.end = (0, 0)
        self.top = self.left = 0
        self.build_tabs()
        self.refresh()

    def new(self):
        self.finish_edit()
        if self.confirm_discard():
            self.load_wb(Workbook())

    def open(self):
        self.finish_edit()
        if not self.confirm_discard():
            return
        path = filedialog.askopenfilename(parent=self.root, title='Abrir',
                                          filetypes=[('Pasta de trabalho do Excel', '*.xlsx'), ('Todos os arquivos', '*.*')])
        if not path:
            return
        try:
            wb, rep = load_xlsx(path)
        except Exception as ex:
            messagebox.showerror('Abrir', f'Não foi possível abrir o arquivo:\n{ex}', parent=self.root)
            return
        self.load_wb(wb)
        self.show_report(rep)

    def show_report(self, rep):
        keep = [x for x in rep if not x.endswith(LOST)]
        lost = [x[:-len(LOST)] for x in rep if x.endswith(LOST)]
        msg = ''
        if keep:
            msg += ('Recursos preservados no arquivo ao salvar, mas que não aparecem aqui:\n\n• '
                    + '\n• '.join(keep) + '\n\n')
        if lost:
            msg += 'Recursos que serão PERDIDOS se você salvar:\n\n• ' + '\n• '.join(lost)
        if msg:
            messagebox.showwarning('Relatório de compatibilidade', msg.strip(), parent=self.root)

    def save(self):
        self.finish_edit()
        if not self.wb.path:
            return self.save_as()
        had_src = self.wb.src_path
        try:
            based = save_xlsx(self.wb, self.wb.path)
        except Exception as ex:
            messagebox.showerror('Salvar', f'Não foi possível salvar:\n{ex}\n\n'
                                 'Se o arquivo estiver aberto em outro programa, feche-o e tente de novo.', parent=self.root)
            return False
        self.wb.dirty = False
        self.refresh()
        if had_src and not based:
            messagebox.showwarning('Salvar', 'O arquivo original não pôde ser lido, então a planilha foi gravada do zero: '
                                   'recursos que não aparecem aqui (bordas, mesclagens, filtros...) não foram mantidos.',
                                   parent=self.root)
        return True

    def save_as(self):
        self.finish_edit()
        path = filedialog.asksaveasfilename(parent=self.root, title='Salvar como', defaultextension='.xlsx',
                                            filetypes=[('Pasta de trabalho do Excel', '*.xlsx')])
        if not path:
            return False
        self.wb.path = path
        return self.save()

    def on_close(self):
        self.finish_edit(refocus=False)
        if self.confirm_discard():
            self.root.destroy()


def main():
    base = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
    ico = os.path.join(base, 'planilha.ico')
    if sys.platform == 'win32':
        try:  # ícone próprio na barra de tarefas (não agrupa sob python.exe)
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID('Planilha.App')
        except Exception:
            pass
    root = tk.Tk()
    try:
        root.iconbitmap(default=ico)
    except (tk.TclError, OSError):
        pass
    app = App(root)
    if len(sys.argv) > 1:
        wb, rep = load_xlsx(sys.argv[1])
        app.load_wb(wb)
        app.show_report(rep)
    root.mainloop()


if __name__ == '__main__':
    main()
