"""Testes da Planilha. Rodar: python test_planilha.py"""
import os
import tempfile
from datetime import datetime

import openpyxl
from openpyxl.styles import PatternFill

from planilha import (DEF, LOST, Date, Err, Workbook, fmt_number, load_xlsx, nf_code, parse_nf, parse_tsv,
                      px_to_width, ref_rc, save_xlsx, shift_refs, to_excel_formula, to_number)


def put(wb, sh, **cells):
    for name, raw in cells.items():
        wb.set_cell(sh, *ref_rc(name), raw)


def val(sh, name):
    return sh.value(*ref_rc(name))


def calc(formula, **cells):
    wb = Workbook()
    sh = wb.sheets[0]
    put(wb, sh, **cells, Z99=formula)
    return val(sh, 'Z99')


DATA = dict(A1='1', A2='2', A3='3', B1='10', B2='20')


def test_functions():
    for f in ('=SOMA(A1:A3)', '=SUM(A1:A3)', '=soma(a1:a3)', '=SOMA(A1;A2;A3)', '=SUM(A1,A2,A3)'):
        assert calc(f, **DATA) == 6, f
    for f in ('=MÉDIA(A1:A3)', '=MEDIA(A1:A3)', '=AVERAGE(A1:A3)', '=média(A1:A3)'):
        assert calc(f, **DATA) == 2, f
    for f in ('=MÍNIMO(A1:B2)', '=MINIMO(A1:B2)', '=MIN(A1:B2)'):
        assert calc(f, **DATA) == 1, f
    for f in ('=MÁXIMO(A1:B2)', '=MAXIMO(A1:B2)', '=MAX(A1:B2)'):
        assert calc(f, **DATA) == 20, f
    assert calc('=SOMA(A1:A3;B1:B2;100)', **DATA) == 136
    assert calc('=SOMA(A3:A1)', **DATA) == 6               # intervalo invertido
    assert calc('=MÉDIA(A1:A10)', **DATA) == 2             # vazias ignoradas
    assert calc('=SOMA(A1:A3)', A1='1', A2='texto', A3='2') == 3  # texto ignorado
    assert calc('=MAX(SOMA(A1:A2);A3*10)', **DATA) == 30
    assert calc('=SOMA(A1:A3)*2', **DATA) == 12


def test_arithmetic():
    assert calc('=2+3*4') == 14
    assert calc('=(2+3)*4') == 20
    assert calc('=10-4-3') == 3
    assert calc('=12/3/2') == 2
    assert calc('=-A1+B1', **DATA) == 9
    assert calc('=--2') == 2
    assert calc('=1.5*2') == 3
    assert calc('=A1+Z1', **DATA) == 1                     # vazia vale 0
    assert calc('= A1 * ( B1 + B2 ) ', **DATA) == 30
    assert calc('=$A$1+A$2', **DATA) == 3


def test_errors_and_circular():
    assert calc('=1/0') == '#DIV/0!' and isinstance(calc('=1/0'), Err)
    assert calc('=A1/Z1', **DATA) == '#DIV/0!'
    assert calc('=MÉDIA(Z1:Z5)') == '#DIV/0!'
    for bad in ('=SOMA(', '=2+', '=FOO(1)', '=A1:A3', '=A1+texto', '=1 2', '=)', '=A0'):
        assert calc(bad, **DATA) == '#ERRO!', bad
    assert calc('=A1+1', A1='abc') == '#ERRO!'
    wb = Workbook()
    sh = wb.sheets[0]
    put(wb, sh, A1='=B1', B1='=A1', C1='=C1', D1='=A1+1', E1='=1/0', F1='=E1*2')
    assert val(sh, 'A1') == '#CIRC!' and val(sh, 'B1') == '#CIRC!' and val(sh, 'C1') == '#CIRC!'
    assert val(sh, 'D1') == '#CIRC!'
    assert val(sh, 'F1') == '#DIV/0!'                     # erro propaga
    put(wb, sh, B1='5')                                    # quebra o ciclo
    assert val(sh, 'A1') == 5 and val(sh, 'D1') == 6


def test_recompute():
    wb = Workbook()
    sh = wb.sheets[0]
    put(wb, sh, A1='2', B1='=A1*2', C1='=B1+A1', D1='=SOMA(A1:C1)')
    assert (val(sh, 'B1'), val(sh, 'C1'), val(sh, 'D1')) == (4, 6, 12)
    put(wb, sh, A1='10')
    assert (val(sh, 'B1'), val(sh, 'C1'), val(sh, 'D1')) == (20, 30, 60)
    put(wb, sh, A1='')
    assert val(sh, 'D1') == 0


def test_formatting():
    assert fmt_number(1234.56, 'Número') == '1.234,56'
    assert fmt_number(1234567.891, 'Número') == '1.234.567,89'
    assert fmt_number(1234.56, 'Moeda') == 'R$ 1.234,56'
    assert fmt_number(-5, 'Moeda') == '-R$ 5,00'
    assert fmt_number(0.125, 'Porcentagem') == '12,50%'
    assert fmt_number(3.0) == '3' and fmt_number(2.5) == '2,5'
    assert to_number('1,5') == 1.5 and to_number('1.234,56') == 1234.56 and to_number('7') == 7
    assert to_number('abc') is None and to_number('inf') is None
    wb = Workbook()
    sh = wb.sheets[0]
    put(wb, sh, A1='1234,5')
    wb.set_format(sh, 0, 0, 0, 0, nf='Moeda')
    assert sh.text(0, 0) == 'R$ 1.234,50'


def test_undo_redo():
    wb = Workbook()
    sh = wb.sheets[0]
    wb.set_cell(sh, 0, 0, '1')
    wb.set_cell(sh, 0, 0, '2')
    wb.set_format(sh, 0, 0, 1, 1, bold=True)
    assert sh.getfmt(1, 1)[0] and sh.raw(0, 0) == '2'
    wb.undo()
    assert sh.getfmt(0, 0) == DEF and sh.raw(0, 0) == '2'
    wb.undo()
    assert sh.raw(0, 0) == '1'
    wb.undo()
    assert sh.raw(0, 0) == '' and wb.undo() is None
    wb.redo()
    assert sh.raw(0, 0) == '1'
    wb.redo(), wb.redo()
    assert sh.raw(0, 0) == '2' and sh.getfmt(0, 0)[0] and wb.redo() is None
    wb.undo()
    wb.set_cell(sh, 5, 5, 'x')                            # nova ação limpa o refazer
    assert wb.redo() is None
    put(wb, sh, A1='5', B1='=A1*2')
    wb.clear(sh, 0, 0, 0, 1)
    assert val(sh, 'B1') is None
    wb.undo()
    assert val(sh, 'B1') == 10                             # cache recalculado após desfazer


def test_copy_paste():
    wb = Workbook()
    sh = wb.sheets[0]
    put(wb, sh, A1='1', A2='2', B1='=A1*2', B2='=$A$2*2')
    wb.set_format(sh, 0, 1, 0, 1, bold=True)
    block = sh.copy_block(0, 0, 1, 1)
    end = wb.paste_block(sh, 3, 2, block, origin=(0, 0))  # para C4
    assert end == (4, 3)
    assert sh.raw(3, 2) == '1' and sh.raw(3, 3) == '=C4*2' and sh.raw(4, 3) == '=$A$2*2'
    assert sh.getfmt(3, 3)[0] and val(sh, 'D4') == 2 and val(sh, 'D5') == 4
    wb.undo()
    assert sh.raw(3, 3) == ''
    assert shift_refs('=SOMA(A1:B2)+$C$1', 1, 1) == '=SOMA(B2:C3)+$C$1'
    assert shift_refs('=A1', -1, 0) == '=#REF!'
    # TSV (ex.: copiado do Excel)
    assert sh.to_tsv(0, 0, 1, 1) == '1\t2\n2\t4'
    rows = parse_tsv('a\t1,5\r\nb\t2\r\n')
    assert rows == [['a', '1,5'], ['b', '2']]
    wb.paste_block(sh, 10, 0, [[(v, None) for v in row] for row in rows])
    assert val(sh, 'B11') == 1.5 and val(sh, 'A12') == 'b'
    # recortar = copiar + limpar (desfazível)
    wb.clear(sh, 0, 0, 1, 1, formats=True)
    assert sh.raw(0, 1) == '' and sh.getfmt(0, 1) == DEF
    wb.undo()
    assert sh.raw(0, 1) == '=A1*2' and sh.getfmt(0, 1)[0]


def test_xlsx_roundtrip():
    wb = Workbook()
    s1 = wb.sheets[0]
    wb.rename_sheet(0, 'Dados')
    s2 = wb.add_sheet('Resumo')
    put(wb, s1, A1='1', A2='2,5', A3='3', A4='=SOMA(A1:A3)', A5='=média(A1;A3)', B1='Olá, ação', B2="'123",
        C1='=MÍNIMO(A1:A3)*2+MÁXIMO(A1:A3)')
    wb.set_format(s1, 0, 0, 0, 0, bold=True)
    wb.set_format(s1, 1, 0, 1, 0, italic=True, nf='Número')
    wb.set_format(s1, 2, 0, 2, 0, nf='Moeda')
    wb.set_format(s1, 3, 0, 3, 0, nf='Porcentagem', bold=True, italic=True)
    put(wb, s2, A1='=2*3')
    path = os.path.join(tempfile.mkdtemp(), 'teste.xlsx')
    save_xlsx(wb, path)

    x = openpyxl.load_workbook(path)
    assert x.sheetnames == ['Dados', 'Resumo']
    ws = x['Dados']
    assert ws['A4'].value == '=SUM(A1:A3)'
    assert ws['A5'].value == '=AVERAGE(A1,A3)'
    assert ws['C1'].value == '=MIN(A1:A3)*2+MAX(A1:A3)'
    assert ws['A1'].value == 1 and ws['A2'].value == 2.5 and ws['B2'].value == '123'
    assert ws['A1'].font.b and ws['A2'].font.i and not ws['A2'].font.b
    assert ws['A2'].number_format == '#,##0.00'
    assert ws['A3'].number_format == '"R$" #,##0.00'
    assert ws['A4'].number_format == '0.00%'

    wb2, rep = load_xlsx(path)
    assert rep == [], rep
    assert [s.name for s in wb2.sheets] == ['Dados', 'Resumo']
    t1, t2 = wb2.sheets
    assert t1.raw(3, 0) == '=SUM(A1:A3)' and val(t1, 'A4') == 6.5
    assert val(t1, 'A5') == 2 and val(t1, 'C1') == 5 and val(t1, 'B1') == 'Olá, ação' and val(t1, 'B2') == '123'
    assert val(t1, 'A2') == 2.5 and val(t2, 'A1') == 6
    for (r, c), f in s1.fmt.items():
        assert t1.getfmt(r, c) == f, (r, c)
    assert t1.text(3, 0) == '650,00%' and t1.text(2, 0) == 'R$ 3,00' and t1.text(1, 0) == '2,50'
    assert not wb2.dirty


def test_compat_report():
    x = openpyxl.Workbook()
    ws = x.active
    ws.title = 'Plan'
    ws['A1'] = 5
    ws['A2'] = '=VLOOKUP(A1,B1:C3,2,FALSE)'
    ws['A3'] = '=Outra!A1+1'
    ws['A4'] = '=sum(A1:A1)'
    ws['B1'] = 7
    ws['B1'].fill = PatternFill('solid', fgColor='FFFF00')
    ws['B2'] = 0.5
    ws['B2'].number_format = '0.00E+00'
    ws.merge_cells('D1:E2')
    path = os.path.join(tempfile.mkdtemp(), 'compat.xlsx')
    x.save(path)
    wb, rep = load_xlsx(path)
    text = '\n'.join(rep)
    for word in ('Células mescladas', 'Formato de número "0.00E+00"', 'VLOOKUP', '2 fórmula(s)'):
        assert word in text, (word, rep)
    assert 'Cores' not in text and not any(r.endswith(LOST) for r in rep), rep  # cores agora são exibidas
    sh = wb.sheets[0]
    assert sh.getfmt(0, 1).fill == 'FFFF00' and sh.getfmt(1, 1).code == '0.00E+00'
    assert sh.raw(1, 0) == "'=VLOOKUP(A1,B1:C3,2,FALSE)" and val(sh, 'A2') == '=VLOOKUP(A1,B1:C3,2,FALSE)'
    assert val(sh, 'A4') == 5
    save_xlsx(wb, path)                                    # fórmula não suportada volta como fórmula
    assert openpyxl.load_workbook(path)['Plan']['A2'].value == '=VLOOKUP(A1,B1:C3,2,FALSE)'
    assert to_excel_formula('=máximo(A1;soma(B1:B2))') == '=MAX(A1,SUM(B1:B2))'


def test_sheets():
    wb = Workbook()
    wb.add_sheet()
    assert [s.name for s in wb.sheets] == ['Planilha1', 'Planilha2']
    for bad in ('', 'planilha1', 'a/b', 'x' * 32):
        try:
            wb.rename_sheet(1, bad)
            raise AssertionError(bad)
        except ValueError:
            pass
    wb.delete_sheet(1)
    try:
        wb.delete_sheet(0)
        raise AssertionError('apagou a última')
    except ValueError:
        pass


def tmp(name):
    return os.path.join(tempfile.mkdtemp(), name)


def test_dates():
    wb = Workbook()
    sh = wb.sheets[0]
    put(wb, sh, A1='26/01/1991', A2='26/01/1991 08:30', A3='31/02/2020', A4='=MÁXIMO(A1:A2)', A5='=A1+1')
    assert isinstance(val(sh, 'A1'), Date) and sh.text(0, 0) == '26/01/1991'
    assert sh.text(1, 0) == '26/01/1991 08:30'
    assert val(sh, 'A3') == '31/02/2020'                   # data inválida fica texto
    assert val(sh, 'A4') == val(sh, 'A2') and sh.text(3, 0) == '26/01/1991 08:30'  # MÁXIMO de datas é data
    assert sh.text(4, 0) == '33265'                        # conta com data dá número de série
    wb.set_format(sh, 4, 0, 4, 0, nf='Data')
    assert sh.text(4, 0) == '27/01/1991'
    path = tmp('datas.xlsx')
    save_xlsx(wb, path)
    ws = openpyxl.load_workbook(path).active
    assert ws['A1'].value == datetime(1991, 1, 26) and ws['A1'].number_format == 'dd/mm/yyyy'
    assert ws['A2'].value == datetime(1991, 1, 26, 8, 30) and ws['A2'].number_format == 'dd/mm/yyyy hh:mm'
    assert ws['A3'].value == '31/02/2020'
    # arquivo com datetime (como o do RH): não pode virar "1991-01-26 00:00:00"
    x = openpyxl.Workbook()
    x.active['A1'] = datetime(1991, 1, 26)
    x.active['A2'] = '26/01/1991'                          # texto no Excel continua texto
    x.save(path)
    wb2, rep = load_xlsx(path)
    s2 = wb2.sheets[0]
    assert s2.text(0, 0) == '26/01/1991' and isinstance(val(s2, 'A1'), Date), (s2.raw(0, 0), rep)
    assert s2.getfmt(0, 0).nf == 'Data' and val(s2, 'A2') == '26/01/1991' and rep == [], rep


def test_number_formats():
    cases = {'0%': ('Porcentagem', 0), '0.0%': ('Porcentagem', 1), '0.00%': ('Porcentagem', 2),
             '#,##0': ('Número', 0), '#,##0.00': ('Número', 2), '0.00': ('Número', 2),
             '"R$" #,##0.00': ('Moeda', 2), '\\R\\$\\ #,##0.00': ('Moeda', 2), '[$R$-416] #,##0.00': ('Moeda', 2),
             '"R$"\\ #,##0.00;[Red]\\-"R$"\\ #,##0.00': ('Moeda', 2),
             '_-"R$"\\ * #,##0.00_-;\\-"R$"\\ * #,##0.00_-;_-"R$"\\ * "-"??_-;_-@_-': ('Moeda', 2),
             'dd/mm/yyyy': ('Data', 2), 'yyyy-mm-dd': ('Data', 2), 'mm-dd-yy': ('Data', 2), 'd/m/yy': ('Data', 2),
             '[$-416]dd/mm/yyyy hh:mm': ('Data', 2), 'General': ('Geral', 2)}
    for code, want in cases.items():
        assert parse_nf(code) == want, (code, parse_nf(code))
    for bad in ('0.00E+00', '@', 'h:mm', '# ?/?'):
        assert parse_nf(bad) is None, bad
    shown = {'0.0%': '12,3%', '\\R\\$\\ #,##0.00': 'R$ 0,12', '#,##0': '0', '0%': '12%'}
    x = openpyxl.Workbook()
    ws = x.active
    for i, code in enumerate(list(shown) + ['0.00E+00'], 1):
        ws.cell(i, 1, 0.1234).number_format = code
    path = tmp('formatos.xlsx')
    x.save(path)
    wb, rep = load_xlsx(path)
    sh = wb.sheets[0]
    for i, (code, txt) in enumerate(shown.items()):
        assert sh.text(i, 0) == txt, (code, sh.text(i, 0))
    assert sh.text(4, 0) == '0,1234' and rep == ['Formato de número "0.00E+00" (exibido como Geral)'], rep
    # casas decimais + / - (desfazível); só a célula alterada perde o código original
    wb.set_format(sh, 0, 0, 0, 0, dec_delta=1)
    assert sh.text(0, 0) == '12,34%' and sh.getfmt(0, 0).code is None
    for _ in range(9):
        wb.set_format(sh, 0, 0, 0, 0, dec_delta=-1)
    assert sh.text(0, 0) == '12%' and sh.getfmt(0, 0).dec == 0
    wb.set_format(sh, 9, 0, 9, 0, dec_delta=1)              # Geral -> Número
    assert sh.getfmt(9, 0)[2:4] == ('Número', 3)
    wb.undo()
    assert sh.getfmt(9, 0) == DEF
    wb.set_cell(sh, 1, 0, '2')                              # muda o valor, mantém o formato
    save_xlsx(wb, path)
    ws = openpyxl.load_workbook(path).active
    assert ws['A1'].number_format == '0%' and ws['A2'].number_format == '\\R\\$\\ #,##0.00'
    assert ws['A2'].value == 2 and ws['A3'].number_format == '#,##0' and ws['A5'].number_format == '0.00E+00'
    assert nf_code('Moeda', 0) == '"R$" #,##0' and fmt_number(1234.6, 'Número', 0) == '1.235'


def test_preservation_roundtrip():
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Border, Font, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    x = openpyxl.Workbook()
    ws = x.active
    ws.title = 'Pessoas'
    for row in [['NOME', 'SALARIO'], ['Fulano Exemplo', 1000], ['Beltrana Exemplo', 2000], ['Ciclano', 3000]]:
        ws.append(row)
    ws.merge_cells('D1:E1')
    ws['D1'] = 'mesclado'
    ws['A2'].border = Border(left=Side(style='thin'), bottom=Side(style='thin'))
    ws['A3'].font = Font(color=openpyxl.styles.colors.Color(theme=4), b=True)
    ws.auto_filter.ref = 'A1:B4'
    ws.conditional_formatting.add('B2:B4', CellIsRule(operator='greaterThan', formula=['1500'], fill=PatternFill('solid', fgColor='FF0000')))
    dv = DataValidation(type='list', formula1='"S,N"')
    dv.add('C2:C4')
    ws.add_data_validation(dv)
    ws.freeze_panes = 'A2'
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions.group('F', 'H', hidden=False)      # faixa <col min=6 max=8>
    ws.column_dimensions['F'].width = 20
    x.create_sheet('Velha')['A1'] = 'renomear'
    x.create_sheet('Apagar')['A1'] = 'x'
    src = tmp('origem.xlsx')
    x.save(src)

    wb, rep = load_xlsx(src)
    text = '\n'.join(rep)
    for word in ('Células mescladas', 'Bordas', 'Filtros', 'Formatação condicional', 'Validação de dados'):
        assert word in text, (word, rep)
    sh = wb.sheets[0]
    assert sh.freeze == (1, 0) and sh.colw[0] == 215 and sh.colw[5] == sh.colw[7] == 145
    wb.set_cell(sh, 1, 1, '1234,5')                         # edita B2
    wb.set_cell(sh, 3, 0, '')                               # esvazia A4
    wb.set_format(sh, 0, 0, 0, 0, italic=True)              # A1: negrito? não; itálico sim
    sh.colw[6] = 50                                          # G, dentro da faixa F:H
    wb.rename_sheet(1, 'Nova')
    wb.delete_sheet(2)
    wb.add_sheet('Extra')
    wb.set_cell(wb.sheets[2], 0, 0, '=2*3')
    out = tmp('salvo_como.xlsx')                             # "Salvar como" usa a origem como base
    assert save_xlsx(wb, out) is True

    y = openpyxl.load_workbook(out)
    assert y.sheetnames == ['Pessoas', 'Nova', 'Extra']
    ws = y['Pessoas']
    assert [str(r) for r in ws.merged_cells.ranges] == ['D1:E1'] and ws['D1'].value == 'mesclado'
    assert ws['A2'].border.left.style == 'thin' and ws['A2'].border.bottom.style == 'thin'
    assert ws['A3'].font.color.theme == 4 and ws['A3'].font.b
    assert ws.auto_filter.ref == 'A1:B4' and len(list(ws.conditional_formatting)) == 1
    assert len(ws.data_validations.dataValidation) == 1 and ws.freeze_panes == 'A2'
    assert ws.column_dimensions['A'].width == 30 and ws['B2'].value == 1234.5 and ws['A4'].value is None
    assert ws['A1'].font.i and ws['A1'].value == 'NOME' and ws['A2'].value == 'Fulano Exemplo'
    widths = {c: d.width for c, d in ws.column_dimensions.items()}
    assert widths['F'] == 20 and widths['G'] == px_to_width(50) and widths['H'] == 20, widths
    assert all(d.min == d.max for k, d in ws.column_dimensions.items() if k in 'FGH')  # sem faixas sobrepostas
    assert y['Nova']['A1'].value == 'renomear' and y['Extra']['A1'].value == '=2*3'
    # o arquivo salvo vira a nova base; sem base legível, grava do zero
    assert wb.src_path == out
    wb.src_path = tmp('sumiu.xlsx')
    assert save_xlsx(wb, out) is False and openpyxl.load_workbook(out)['Pessoas']['B2'].value == 1234.5


def test_colors_alignment():
    wb = Workbook()
    sh = wb.sheets[0]
    put(wb, sh, A1='texto longo que quebra', B1='5')
    wb.set_format(sh, 0, 0, 0, 0, color='FF0000', fill='FFFF00', align='center', wrap=True)
    wb.set_format(sh, 0, 1, 0, 1, align='right', color='0000FF')
    f = sh.getfmt(0, 0)
    assert (f.color, f.fill, f.align, f.wrap) == ('FF0000', 'FFFF00', 'center', True)
    wb.set_format(sh, 0, 1, 0, 1, color=None)               # "sem cor"
    assert sh.getfmt(0, 1).color is None
    wb.undo()
    assert sh.getfmt(0, 1).color == '0000FF'
    wb.undo()
    assert sh.getfmt(0, 1) == DEF
    wb.redo()
    path = tmp('cores.xlsx')
    save_xlsx(wb, path)
    ws = openpyxl.load_workbook(path).active
    assert ws['A1'].font.color.rgb == 'FFFF0000' and ws['A1'].fill.fgColor.rgb == 'FFFFFF00'
    assert ws['A1'].alignment.horizontal == 'center' and ws['A1'].alignment.wrap_text
    assert ws['B1'].alignment.horizontal == 'right' and ws['B1'].font.color.rgb == 'FF0000FF'
    wb2, rep = load_xlsx(path)
    assert rep == [] and wb2.sheets[0].fmt == sh.fmt, (rep, wb2.sheets[0].fmt, sh.fmt)
    # cor indexada (paleta antiga do Excel)
    from openpyxl.styles import Font
    x = openpyxl.Workbook()
    x.active['A1'] = 'x'
    x.active['A1'].font = Font(color=openpyxl.styles.colors.Color(indexed=10))
    x.save(path)
    assert load_xlsx(path)[0].sheets[0].getfmt(0, 0).color == 'FF0000'


def test_ui_smoke():
    import tkinter as tk
    from planilha import App
    try:
        root = tk.Tk()
    except tk.TclError as ex:
        print('  (sem display, smoke da UI pulado:', ex, ')')
        return
    try:
        root.withdraw()
        app = App(root)
        root.update()
        app.start_edit('=2*21')
        app.finish_edit(1, 0)
        assert app.sh.value(0, 0) == 42 and app.cur == (1, 0)
        assert root.title() == 'Sem título* - planilha_code'
        app.cur, app.end = (0, 0), (0, 0)
        app.copy()
        app.cur = app.end = (0, 1)
        app.paste()
        assert app.sh.raw(0, 1) == '=2*21'
        app.toggle_bold()
        assert app.sh.getfmt(0, 1)[0]
        app.undo()
        assert not app.sh.getfmt(0, 1)[0]
        app.add_sheet()
        assert app.si == 1 and len(app.tabbtns) == 2
        app.move(150, 30)                                  # cresce a grade
        root.update()
        # painéis congelados + larguras/alturas variáveis + cores/quebra
        wb = Workbook()
        sh = wb.sheets[0]
        sh.freeze, sh.colw, sh.rowh = (1, 1), {0: 200, 1: 40, 3: 150}, {0: 44}
        for r in range(300):
            wb.set_cell(sh, r, 0, f'Linha {r} com um texto comprido para quebrar')
        wb.set_format(sh, 0, 0, 0, 3, bold=True, fill='DDEEFF', wrap=True, align='center')
        app.load_wb(wb)
        root.update()
        app.move(250, 20)
        root.update()
        assert app.top > 1 and app.left > 1
        cols = app.axis(True)
        assert cols[0][:3] == (0, 50, 200) and cols[1][0] == app.left  # coluna A fixa
        assert app.axis(False)[0] == (0, 22, 44)
        app.autofit(0)
        assert sh.colw[0] > 200
        app.freeze_panes()
        assert sh.freeze == (250, 20)
        app.unfreeze()
        assert sh.freeze == (0, 0)
        root.update()
    finally:
        root.destroy()


if __name__ == '__main__':
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith('test_') and callable(f)]
    for name, fn in tests:
        fn()
        print('ok ', name)
    print(f'\n{len(tests)} testes passaram.')
